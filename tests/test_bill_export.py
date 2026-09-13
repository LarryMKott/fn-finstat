"""流水导出测试：xlsx / csv 格式、筛选条件与账号隔离"""

import csv
import io

from openpyxl import load_workbook

from app.db.dao.bill_dao import BillDAO
from tests.conftest import USER_A, USER_B, make_bill_records

A_HEADERS = {"X-Trim-Userid": USER_A}
B_HEADERS = {"X-Trim-Userid": USER_B}


def test_export_xlsx_all_columns(client):
    BillDAO.insert_many(
        make_bill_records(
            2,
            prefix="EX",
            tx_time="2026-08-01 10:00:00",
            category="餐饮",
            tags="出差",
            user_id="",
        ),
        USER_A,
    )
    res = client.get("/api/bill/export?format=xlsx", headers=A_HEADERS)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:6] == (
        "交易时间",
        "账户",
        "类型",
        "商户/交易对方",
        "金额(元)",
        "分类",
    )
    assert len(rows) == 3  # 表头 + 2 条
    assert rows[1][2] == "支出"
    assert rows[1][6] == "出差"


def test_export_csv_filters_and_scoping(client):
    BillDAO.insert_many(
        make_bill_records(
            2, prefix="EC1", tx_time="2026-07-01 10:00:00", account="wechat"
        ),
        USER_A,
    )
    BillDAO.insert_many(
        make_bill_records(
            1, prefix="EC2", tx_time="2026-07-02 10:00:00", account="alipay"
        ),
        USER_A,
    )
    BillDAO.insert_many(
        make_bill_records(
            1, prefix="EC3", tx_time="2026-07-03 10:00:00", account="wechat"
        ),
        USER_B,
    )

    res = client.get("/api/bill/export?format=csv&account=alipay", headers=A_HEADERS)
    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert len(rows) == 2  # 表头 + 仅 alipay 那 1 条（B 账号数据不出现）
    assert rows[1][8] == "EC2-0000"  # 交易单号列（prefix-序号）

    # 回收站流水不导出
    bill_id = client.get("/api/bill/list", headers=A_HEADERS).json()["items"][0]["id"]
    client.delete(f"/api/bill/{bill_id}", headers=A_HEADERS)
    res = client.get("/api/bill/export?format=csv", headers=A_HEADERS)
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
    assert len(rows) == 3  # 2 条未删除 + 表头


def test_export_invalid_format_rejected(client):
    res = client.get("/api/bill/export?format=xls", headers=A_HEADERS)
    assert res.status_code == 400
