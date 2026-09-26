"""导出文件构建：流水行 → xlsx / CSV 字节流（导出列与解析导入字段对应）

公式注入防护（安全审计 M8-7）：

- **CSV**：Excel 会把 `=` `+` `-` `@` 开头的单元格当公式求值 → `csv_safe` 前缀单引号。
  天真的 `startswith` 判定可被「前导空白/控制字符」绕过（`" =cmd"`、`"\t@SUM(1)"`），
  故先剥离这些字符再判定。
- **xlsx**：openpyxl 写字符串时**只对 `=` 开头自动升级为公式**（实测 `data_type='f'`），
  `+` / `-` / `@` 仍写为文本（`'s'`）→ 文本列遇 `=` 前缀需显式钉死 `data_type='s'`，
  既拦注入又不改动用户数据（对比 CSV 的加引号方案，xlsx 无需改内容）。
"""

import csv
import io

from openpyxl import Workbook

from app.core.constants import ACCOUNT_LABELS, TX_TYPE_LABELS

# 导出列定义（表头 → 取值键），xlsx 与 CSV 共用同一份顺序
EXPORT_COLUMNS = [
    ("交易时间", "tx_time"),
    ("账户", "account"),
    ("类型", "tx_type_label"),
    ("商户/交易对方", "merchant"),
    ("金额(元)", "amount"),
    ("分类", "category"),
    ("标签", "tags"),
    ("报销", "reimbursed_label"),
    ("交易单号", "tx_id"),
    ("备注", "remark"),
]

# 金额列下标：CSV 侧保持数值原样，避免破坏 Excel 二次统计
_AMOUNT_COL = 4

# CSV 侧 Excel 会按公式求值的前缀（openpyxl 侧只有 "=" 会被升级为公式）
_CSV_DANGEROUS_PREFIXES = ("=", "+", "-", "@")

# 需先剥离再判定的前导字符：不可见控制字符 + 各类空白
# （C0 控制符除 \t\n\r 外、NUL、以及 Unicode 空白）
_STRIP_FOR_PREFIX_CHECK = (
    "\x00\x01\x02\x03\x04\x05\x06\x07\x08\x0b\x0c\x0e\x0f"
    "\x10\x11\x12\x13\x14\x15\x16\x17\x18\x19\x1a\x1b\x1c\x1d\x1e\x1f"
    " \t\r\n\x0b\x0c\xa0\u1680\u2000\u2001\u2002\u2003\u2004\u2005\u2006"
    "\u2007\u2008\u2009\u200a\u2028\u2029\u202f\u205f\u3000\ufeff"
)


def _lstrip_invisible(text: str) -> str:
    """剥掉开头的空白与不可见控制字符（用于前缀判定，不改动原值）"""
    return text.lstrip(_STRIP_FOR_PREFIX_CHECK)


def _export_rows(bills: list[dict]) -> list[list]:
    """把流水字典转为按 EXPORT_COLUMNS 顺序排列的二维表（首行为表头）"""
    header = [c[0] for c in EXPORT_COLUMNS]
    rows = [header]
    for b in bills:
        account = ACCOUNT_LABELS.get(b.get("account"), b.get("account"))
        rows.append(
            [
                b.get("tx_time", ""),
                account,
                TX_TYPE_LABELS.get(b.get("tx_type"), b.get("tx_type", "")),
                b.get("merchant", ""),
                b.get("amount", 0),
                b.get("category", ""),
                b.get("tags", ""),
                "是" if b.get("reimbursed") else "",
                b.get("tx_id") or "",
                b.get("remark", ""),
            ]
        )
    return rows


def build_xlsx(bills: list[dict]) -> bytes:
    """构建 xlsx（openpyxl 内存构建）；金额写为数值便于 Excel 二次统计

    文本列防公式注入：openpyxl 会把 `=` 开头的字符串自动写成公式单元格
    （实测 `data_type='f'`），打开文件即可能触发 `=cmd|...` 这类 DDE 攻击。
    此处对 `=` 开头的文本显式钉死 `data_type='s'`（字符串），原文不动。
    `+` / `-` / `@` 前缀 openpyxl 本就写为文本，无需处理（与实测一致）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "流水"
    rows = _export_rows(bills)
    for row in rows:
        ws.append(row)
    # 钉死被 openpyxl 误升级为公式的文本单元格
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row, start=1):
            if isinstance(cell, str) and _lstrip_invisible(cell).startswith("="):
                ws.cell(row=i, column=j).data_type = "s"
    # 简单设宽：金额/时间列加宽，避免默认宽度挤成 ####
    for idx, width in ((1, 20), (4, 24), (5, 12), (9, 24)):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_safe(value) -> str:
    """防 CSV 公式注入：Excel 打开时 = + - @ 开头的单元格会被按公式求值，
    前缀单引号使其按文本显示

    判定前先剥离前导空白与不可见控制字符 —— 否则 `" =cmd"`、`"\\t@SUM(A1)"`
    会被 Excel 正常解析为公式（实测可绕过朴素 startswith），而原值必须保留
    （只加前缀、不改内容，避免破坏用户数据）。
    """
    text = str(value)
    if _lstrip_invisible(text).startswith(_CSV_DANGEROUS_PREFIXES):
        return "'" + text
    return text


def build_csv(bills: list[dict]) -> bytes:
    """构建 CSV（utf-8-sig 带 BOM，Excel 直接打开不乱码）

    仅 CSV 需要「加单引号」式防护（xlsx 走 data_type 路线，见 build_xlsx）。
    金额列保持数值原样，避免破坏 Excel 二次统计。
    """
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in _export_rows(bills):
        writer.writerow(
            [cell if i == _AMOUNT_COL else csv_safe(cell) for i, cell in enumerate(row)]
        )
    return buf.getvalue().encode("utf-8-sig")
