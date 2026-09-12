"""微信支付账单（.xlsx）解析器

微信导出的 Excel 账单前几行为说明文字，随后是列头行：
    交易时间,交易类型,交易对方,商品,收/支,金额(元),支付方式,当前状态,交易单号,商户单号,备注
"""
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.base import BaseParser
from app.utils.amount import normalize_amount


class WechatParser(BaseParser):
    """微信 xlsx 账单解析器"""

    account = "wechat"

    def parse(self, file_path: Path) -> list[dict]:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        header: list[str] | None = None
        records: list[dict] = []
        try:
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                cells = [str(c).strip() if c is not None else "" for c in row]
                if not any(cells):
                    continue
                if header is None:
                    # 列头行：第一列为“交易时间”
                    if cells and cells[0].startswith("交易时间"):
                        header = cells
                    continue
                record = self._row_to_record(dict(zip(header, cells)))
                if record:
                    records.append(record)
        finally:
            wb.close()
        return records

    def _row_to_record(self, row: dict) -> dict | None:
        """单行转标准流水；缺交易时间或金额的行返回 None"""
        tx_time = row.get("交易时间", "").strip()
        tx_id = row.get("交易单号", "").strip()
        if not tx_time:
            return None

        amount_text = row.get("金额(元)", "").strip()
        amount_text = amount_text.replace("¥", "").replace("￥", "").replace(",", "").strip()
        if not amount_text:
            return None
        try:
            amount = normalize_amount(amount_text)
        except ValueError:
            return None
        if amount <= 0:
            return None

        io = row.get("收/支", "").strip()
        if io == "收入":
            tx_type = "income"
        elif io == "支出":
            tx_type = "expense"
        else:
            # “/”、“不计收支”等转账类流水
            tx_type = "transfer"

        return {
            "tx_time": tx_time,
            "account": self.account,
            "tx_type": tx_type,
            "merchant": row.get("交易对方", "").strip(),
            "amount": amount,
            "category": "",
            "tx_id": tx_id,
            "remark": row.get("备注", "").strip(),
        }
