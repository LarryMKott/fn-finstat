"""账单来源自动识别：读文件头部样例，按各平台表头特征打分判定

NAS 目录导入没有「选择平台」这一步，导入前先识别来源：
- csv：取文件头部样例（多编码探测），定位列头行后按「平台特征列」精确
  匹配计分，得分最高者胜出（特征列取各平台独有/区分度高的列，公共列不计）；
- xlsx：微信账单是唯一 xlsx 来源，按其列头特征确认；
- 内容识别不出时回退文件名关键字（仅允许与后缀匹配的平台，避免误判）。
"""

import csv
import io
import logging
from pathlib import Path

from openpyxl import load_workbook

from app.parsers.alipay_parser import AlipayParser
from app.parsers.base import BaseParser
from app.parsers.csv_common import ENCODINGS
from app.parsers.jd_parser import JdParser
from app.parsers.unionpay_parser import UnionPayParser
from app.parsers.wechat_parser import WechatParser

logger = logging.getLogger(__name__)

# 表头行判定：任一单元格为「交易时间/交易日期」即视为列头行（四个平台通用）
_HEADER_CELLS = ("交易时间", "交易日期")

# 各平台特征列 → 权重（精确匹配表头单元格后求和，最高分且 >0 者胜出）
_CSV_MARKERS: dict[str, dict[str, int]] = {
    "alipay": {
        "交易分类": 3,
        "对方账号": 3,
        "收/付款方式": 3,
        "商家订单号": 2,
        "商品说明": 2,
        "交易订单号": 1,
    },
    "wechat": {
        # 微信 CSV（非官方导出主流但存在）特征列：商户单号/支付方式/当前状态为独有
        "商户单号": 3,
        "支付方式": 2,
        "当前状态": 2,
        "交易单号": 1,
        "交易对方": 1,
    },
    "jd": {
        "商品名称": 3,
        "流水号": 3,
        "交易金额": 1,
        "商户名称": 1,
        "交易单号": 1,
    },
    "unionpay": {
        "交易方式": 3,
        "交易金额(元)": 3,
        "收付款对象": 2,
        "订单号": 1,
    },
}

# 文件名关键字兜底（内容识别失败时使用；wechat 仅 xlsx，其余仅 csv）
_FILENAME_HINTS: dict[str, tuple[str, ...]] = {
    "wechat": ("微信", "wechat", "weixin"),
    "alipay": ("支付宝", "alipay"),
    "jd": ("京东", "jd"),
    "unionpay": ("云闪付", "银联", "unionpay"),
}

# csv 来源识别只读文件头部样例即可覆盖表头行
_SAMPLE_BYTES = 64 * 1024


def build_parser(source: str) -> BaseParser | None:
    """来源 key → 解析器实例；未知来源返回 None"""
    parsers = {
        "alipay": AlipayParser,
        "jd": JdParser,
        "unionpay": UnionPayParser,
        "wechat": WechatParser,
    }
    cls = parsers.get(source)
    return cls() if cls else None


def detect_source(path: Path) -> str:
    """识别账单来源：wechat/alipay/jd/unionpay；无法识别返回空串"""
    ext = path.suffix.lower()
    if ext == ".csv":
        source = _detect_csv(path)
    elif ext == ".xlsx":
        source = _detect_xlsx(path)
    else:
        source = ""
    return source or _detect_by_filename(path.name, ext)


def _detect_by_filename(name: str, ext: str) -> str:
    """文件名关键字兜底；后缀与平台不匹配的关键字不采信"""
    lowered = name.lower()
    for source, hints in _FILENAME_HINTS.items():
        if source == "wechat" and ext != ".xlsx":
            continue
        if source != "wechat" and ext != ".csv":
            continue
        if any(hint in lowered for hint in hints):
            return source
    return ""


def _decode_sample(raw: bytes) -> str:
    """样例字节多编码探测解码；全部失败按 gb18030 替换坏字节兜底"""
    for enc in ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("gb18030", errors="replace")


def _detect_csv(path: Path) -> str:
    """csv：定位列头行后按平台特征列打分，取最高分"""
    try:
        with path.open("rb") as f:
            sample = f.read(_SAMPLE_BYTES)
    except OSError as exc:
        logger.warning("来源识别读取文件失败（%s）：%s", path.name, exc)
        return ""
    reader = csv.reader(io.StringIO(_decode_sample(sample)))
    for row in reader:
        cells = [c.strip() for c in row]
        if not any(cell in _HEADER_CELLS for cell in cells):
            continue
        scores = {
            source: sum(w for col, w in markers.items() if col in cells)
            for source, markers in _CSV_MARKERS.items()
        }
        best = max(scores, key=lambda k: scores[k])
        return best if scores[best] > 0 else ""
    return ""


def _detect_xlsx(path: Path) -> str:
    """xlsx：微信账单是唯一来源，按其列头（交易时间 + 交易单号/商户单号）确认"""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("来源识别打开 xlsx 失败（%s）：%s", path.name, exc)
        return ""
    try:
        ws = wb.active
        for row in ws.iter_rows(values_only=True, max_row=30):
            cells = [str(c).strip() for c in row or ()]
            if not cells or not cells[0].startswith("交易时间"):
                continue
            if "交易单号" in cells or "商户单号" in cells:
                return "wechat"
            return ""
    except Exception as exc:
        logger.warning("来源识别读取 xlsx 失败（%s）：%s", path.name, exc)
        return ""
    finally:
        wb.close()
    return ""
